from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import (
    Alignment,
    Font,
    PatternFill,
)
from openpyxl.utils import (
    get_column_letter,
)


class FinalExporter:

    def __init__(
        self,
        companies_path="data/cache/companies_master.csv",
        web_path="data/cache/web_validated.csv",
        emails_summary_path="data/cache/emails_by_company_clean.csv",
        emails_detail_path="data/cache/emails_found_clean.csv",
        emails_rejected_path="data/cache/emails_rejected.csv",
        search_results_path="data/cache/search_results.csv",
        search_log_path="data/cache/search_run_log.csv",
        crawl_pages_path="data/cache/crawled_pages.csv",
        crawl_log_path="data/cache/crawl_seed_log.csv",
        output_root="data/output",
    ):
        self.companies_path = Path(
            companies_path
        )
        self.web_path = Path(
            web_path
        )

        self.emails_summary_path = Path(
            emails_summary_path
        )

        self.emails_detail_path = Path(
            emails_detail_path
        )

        self.emails_rejected_path = Path(
            emails_rejected_path
        )

        self.search_results_path = Path(
            search_results_path
        )

        self.search_log_path = Path(
            search_log_path
        )

        self.crawl_pages_path = Path(
            crawl_pages_path
        )

        self.crawl_log_path = Path(
            crawl_log_path
        )

        self.output_root = Path(
            output_root
        )

    @staticmethod
    def read(path):
        if not path.exists():
            return pd.DataFrame()

        if path.stat().st_size == 0:
            return pd.DataFrame()

        try:
            return pd.read_csv(
                path,
                dtype=str,
                keep_default_na=False,
            )
        except pd.errors.EmptyDataError:
            return pd.DataFrame()

    def build_final(self):
        companies = self.read(
            self.companies_path
        )

        if companies.empty:
            raise RuntimeError(
                "companies_master.csv está vacío"
            )

        webs = self.read(
            self.web_path
        )

        emails = self.read(
            self.emails_summary_path
        )

        final = companies.copy()

        if not webs.empty:
            web_small = webs.copy()

            web_small = (
                web_small
                .drop_duplicates(
                    subset=[
                        "COMPANY_KEY",
                    ],
                    keep="last",
                )
            )

            web_small = web_small[
                [
                    "COMPANY_KEY",
                    "FINAL_URL",
                    "HTTP_STATUS",
                    "VALIDATION_STATUS",
                ]
            ].rename(
                columns={
                    "FINAL_URL":
                        "WEB",
                    "HTTP_STATUS":
                        "WEB_HTTP_STATUS",
                    "VALIDATION_STATUS":
                        "WEB_ESTADO",
                }
            )

            final = final.merge(
                web_small,
                on="COMPANY_KEY",
                how="left",
            )

        if not emails.empty:
            final = final.merge(
                emails,
                on="COMPANY_KEY",
                how="left",
            )

        for col in final.columns:
            final[col] = (
                final[col]
                .fillna("")
            )

        if "WEB_ESTADO" not in final.columns:
            final[
                "WEB_ESTADO"
            ] = "PENDIENTE"

        final[
            "WEB_ESTADO"
        ] = final[
            "WEB_ESTADO"
        ].replace(
            "",
            "PENDIENTE",
        )

        if "ESTADO_EMAIL" not in final.columns:
            final[
                "ESTADO_EMAIL"
            ] = "SIN_EMAIL"

        final[
            "ESTADO_EMAIL"
        ] = final[
            "ESTADO_EMAIL"
        ].replace(
            "",
            "SIN_EMAIL",
        )

        def process_status(row):
            estado_email = str(
                row.get(
                    "ESTADO_EMAIL",
                    "",
                )
            ).strip()

            if (
                estado_email
                and estado_email != "SIN_EMAIL"
            ):
                return "CON_CONTACTO"

            if (
                row.get(
                    "WEB_ESTADO",
                    "",
                )
                == "OK"
            ):
                return "WEB_SIN_CONTACTO"

            return "PENDIENTE"

        final[
            "ESTADO_PROCESO"
        ] = final.apply(
            process_status,
            axis=1,
        )

        final[
            "FECHA_EXPORTACION"
        ] = datetime.now().isoformat(
            timespec="seconds"
        )

        preferred = [
            "COMPANY_KEY",
            "COMPANY_NAME",
            "NICHE_IDS",
            "ROOT_DOMAIN",
            "WEB",
            "WEB_ESTADO",
            "WEB_HTTP_STATUS",

            "ESTADO_EMAIL",
            "EMAIL_PRINCIPAL",
            "EMAIL_VENTAS",
            "EMAIL_CONTACTO",
            "EMAIL_COMERCIAL",
            "EMAIL_SOPORTE",
            "EMAIL_LOGISTICA",
            "EMAIL_ADMINISTRACION",

            "EMAILS_ROL_INSTITUCIONAL",
            "EMAILS_PERSONALES_PUBLICOS",
            "EMAILS_TODOS",

            "TOTAL_EMAILS",
            "TOTAL_ROL_INSTITUCIONAL",
            "TOTAL_PERSONALES_PUBLICOS",

            "TELEFONOS",

            "BEST_URL",
            "BEST_SCORE",
            "RESULT_COUNT",
            "QUERY_COUNT",
            "PROVIDERS",
            "QUERIES",

            "DISCOVERY_STATUS",
            "DISCOVERED_AT",
            "LAST_SEEN_AT",

            "ESTADO_PROCESO",
            "FECHA_EXPORTACION",
        ]

        ordered = [
            col
            for col in preferred
            if col in final.columns
        ]

        others = [
            col
            for col in final.columns
            if col not in ordered
        ]

        final = final[
            ordered + others
        ]

        return final

    @staticmethod
    def style_workbook(path):
        wb = load_workbook(
            path
        )

        fill = PatternFill(
            "solid",
            fgColor="1F4E78",
        )

        font = Font(
            color="FFFFFF",
            bold=True,
        )

        for ws in wb.worksheets:
            ws.freeze_panes = "A2"

            if ws.max_row >= 1:
                ws.auto_filter.ref = (
                    ws.dimensions
                )

            for cell in ws[1]:
                cell.fill = fill
                cell.font = font
                cell.alignment = (
                    Alignment(
                        horizontal="center",
                        vertical="center",
                    )
                )

            for column_cells in ws.columns:
                max_length = 0

                for cell in column_cells:
                    value = str(
                        cell.value or ""
                    )

                    max_length = max(
                        max_length,
                        len(value),
                    )

                letter = get_column_letter(
                    column_cells[
                        0
                    ].column
                )

                ws.column_dimensions[
                    letter
                ].width = min(
                    max(
                        max_length + 2,
                        12,
                    ),
                    60,
                )

        wb.save(path)

    def run(self):
        final = self.build_final()

        self.output_root.mkdir(
            parents=True,
            exist_ok=True,
        )

        emails_detail = self.read(
            self.emails_detail_path
        )

        rejected = self.read(
            self.emails_rejected_path
        )

        webs = self.read(
            self.web_path
        )

        search_results = self.read(
            self.search_results_path
        )

        search_log = self.read(
            self.search_log_path
        )

        crawl_pages = self.read(
            self.crawl_pages_path
        )

        crawl_log = self.read(
            self.crawl_log_path
        )

        summary = pd.DataFrame(
            [
                {
                    "METRICA":
                        "Empresas descubiertas",
                    "VALOR":
                        len(final),
                },
                {
                    "METRICA":
                        "Webs OK",
                    "VALOR":
                        (
                            final[
                                "WEB_ESTADO"
                            ]
                            .eq("OK")
                            .sum()
                        ),
                },
                {
                    "METRICA":
                        "Empresas con contacto",
                    "VALOR":
                        (
                            final[
                                "ESTADO_PROCESO"
                            ]
                            .eq(
                                "CON_CONTACTO"
                            )
                            .sum()
                        ),
                },
                {
                    "METRICA":
                        "Correos institucionales",
                    "VALOR":
                        (
                            final[
                                "ESTADO_EMAIL"
                            ]
                            .eq(
                                "CON_CORREO_OFICIAL"
                            )
                            .sum()
                        ),
                },
                {
                    "METRICA":
                        "Solo correos personales públicos",
                    "VALOR":
                        (
                            final[
                                "ESTADO_EMAIL"
                            ]
                            .eq(
                                "SOLO_PERSONALES_PUBLICOS"
                            )
                            .sum()
                        ),
                },
                {
                    "METRICA":
                        "Páginas HTML acumuladas",
                    "VALOR":
                        len(
                            crawl_pages
                        ),
                },
            ]
        )

        sheets = {
            "EMPRESAS_FINAL":
                final,

            "EMAILS_DETALLE":
                emails_detail,

            "EMAILS_RECHAZADOS":
                rejected,

            "WEBS_VALIDADAS":
                webs,

            "RESULTADOS_BUSQUEDA":
                search_results,

            "SEARCH_LOG":
                search_log,

            "CRAWL_PAGES":
                crawl_pages,

            "CRAWL_LOG":
                crawl_log,

            "RESUMEN":
                summary,
        }

        latest_path = (
            self.output_root
            / "ZT_DataEmpresasPriv.xlsx"
        )

        dated_path = (
            self.output_root
            / (
                "ZT_DataEmpresasPriv_"
                + datetime.now()
                .strftime(
                    "%Y-%m-%d"
                )
                + ".xlsx"
            )
        )

        csv_path = (
            self.output_root
            / "ZT_DataEmpresasPriv.csv"
        )

        for output_path in [
            latest_path,
            dated_path,
        ]:
            with pd.ExcelWriter(
                output_path,
                engine="openpyxl",
            ) as writer:

                for name, df in sheets.items():

                    df.to_excel(
                        writer,
                        sheet_name=name[:31],
                        index=False,
                    )

            self.style_workbook(
                output_path
            )

        final.to_csv(
            csv_path,
            index=False,
            encoding="utf-8-sig",
        )

        print("=" * 70)
        print(" FINAL EXPORTER")
        print("=" * 70)

        print(
            f"Empresas final         : {len(final):,}"
        )

        print(
            f"Webs OK                : "
            f"{final['WEB_ESTADO'].eq('OK').sum():,}"
        )

        print(
            f"Empresas con contacto  : "
            f"{final['ESTADO_PROCESO'].eq('CON_CONTACTO').sum():,}"
        )

        print()
        print(latest_path)
        print(dated_path)
        print(csv_path)
        print("=" * 70)

        return latest_path
